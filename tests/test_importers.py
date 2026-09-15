"""
tests/test_importers.py
Covers: xlsx_validator, build_db import_roster (FK safety, success path, live-data check)
"""
import os
import sqlite3
import pytest

from data.db import get_connection
from data.migrate import apply_migrations
from core.importers.xlsx_validator import (
    validate_xlsx_structure,
    validate_sqlite_file,
    check_live_data_exists,
)
from core.importers.build_db import import_roster


@pytest.fixture
def migrated_db(tmp_path):
    db_file = str(tmp_path / "school.sqlite")
    conn = get_connection(db_file, pin="123456")
    apply_migrations(conn)
    yield conn
    conn.close()


def _make_minimal_xlsx(path: str, sheet_names: list) -> None:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = sheet_names[0] if sheet_names else "Sheet1"
    for name in sheet_names[1:]:
        wb.create_sheet(name)
    wb.save(path)


# --- xlsx_validator ---

def test_xlsx_validate_missing_file(tmp_path):
    res = validate_xlsx_structure(str(tmp_path / "nonexistent.xlsx"))
    assert res.ok is False
    assert "not found" in res.error.lower()


def test_xlsx_validate_non_excel_file(tmp_path):
    bad = str(tmp_path / "bad.xlsx")
    with open(bad, "w") as f:
        f.write("not a zip")
    res = validate_xlsx_structure(bad)
    assert res.ok is False


def test_xlsx_validate_missing_grade_sheets(tmp_path):
    path = str(tmp_path / "partial.xlsx")
    _make_minimal_xlsx(path, ["ม.1", "ม.2"])
    res = validate_xlsx_structure(path)
    assert res.ok is False


def test_xlsx_validate_all_grade_sheets_ok(tmp_path):
    path = str(tmp_path / "full.xlsx")
    _make_minimal_xlsx(path, ["ม.1", "ม.2", "ม.3", "ม.4", "ม.5", "ม.6"])
    res = validate_xlsx_structure(path)
    assert res.ok is True


def test_validate_sqlite_accepts_migrated_db(migrated_db):
    res = validate_sqlite_file("", conn=migrated_db)
    assert res.ok is True


def test_validate_sqlite_rejects_empty_db(tmp_path):
    empty = str(tmp_path / "empty.sqlite")
    sqlite3.connect(empty).close()
    res = validate_sqlite_file(empty)
    assert res.ok is False


def test_validate_sqlite_rejects_nonexistent(tmp_path):
    res = validate_sqlite_file(str(tmp_path / "no.sqlite"))
    assert res.ok is False


# --- check_live_data_exists ---

def test_live_data_false_on_fresh_db(migrated_db):
    has_live, _ = check_live_data_exists(migrated_db)
    assert has_live is False


def test_live_data_true_when_loans_exist(migrated_db):
    from data.repositories import book_repo, loan_repo
    migrated_db.execute(
        "INSERT INTO students (student_id, full_name, status) VALUES (1, 'ทดสอบ', 'active')"
    )
    migrated_db.commit()
    book_id = book_repo.add_book(migrated_db, title="Test Book")
    copy_ids = book_repo.add_copies(migrated_db, book_id, barcodes=["TEST-01"])
    loan_repo.create_loan(
        migrated_db, copy_id=copy_ids[0], student_id=1,
        borrowed_at="2026-09-01T08:00:00", due_at="2026-09-15T08:00:00"
    )
    has_live, reason = check_live_data_exists(migrated_db)
    assert has_live is True  # either activity_log or loans triggers live-data


# --- import_roster FK safety (BLOCKER 2 regression) ---

def test_import_roster_fk_restored_after_failure(migrated_db, tmp_path, monkeypatch):
    """
    Simulate a failure INSIDE the with conn: block (after PRAGMA FK=OFF).
    FK enforcement must be restored to ON even on exception.
    """
    import openpyxl
    import unittest.mock as mock

    # Build a valid xlsx so load_workbook succeeds
    valid_xlsx = str(tmp_path / "valid.xlsx")
    wb = openpyxl.Workbook()
    for name in ["ม.1", "ม.2", "ม.3", "ม.4", "ม.5", "ม.6"]:
        if wb.active.title == "Sheet":
            wb.active.title = name
        else:
            wb.create_sheet(name)
    wb.save(valid_xlsx)

    # Patch parse_grade_sheet to raise AFTER FK=OFF so the finally block is exercised
    from core.importers import build_db
    original = build_db.parse_grade_sheet
    call_count = [0]
    def boom(ws, cur):
        call_count[0] += 1
        raise RuntimeError("Simulated parse failure after FK=OFF")
    monkeypatch.setattr(build_db, "parse_grade_sheet", boom)

    res = import_roster(valid_xlsx, conn=migrated_db)
    assert res.ok is False
    assert call_count[0] >= 1  # confirm the bomb actually fired

    # FK must be restored to ON
    cur = migrated_db.cursor()
    cur.execute("SELECT * FROM pragma_foreign_keys")
    fk_on = cur.fetchone()[0]
    assert fk_on == 1, "FK enforcement not restored after failed import!"


def test_import_roster_success_with_valid_xlsx(tmp_path, migrated_db):
    import openpyxl
    path = str(tmp_path / "roster.xlsx")
    wb = openpyxl.Workbook()
    grade_sheets = ["ม.1", "ม.2", "ม.3", "ม.4", "ม.5", "ม.6"]
    wb.active.title = grade_sheets[0]
    for name in grade_sheets[1:]:
        wb.create_sheet(name)
    ws = wb[grade_sheets[0]]
    ws["A1"] = "ชั้นมัธยมศึกษาปีที่ 1/1 ปีการศึกษา 2567 ภาคเรียนที่ 1"
    ws["A3"] = "ที่"
    ws["B3"] = "เลขประจำตัว"
    ws["C3"] = "ชื่อ-สกุล"
    ws["A4"] = 1
    ws["B4"] = 10001
    ws["C4"] = "เด็กชายสมชาย ดีมาก"
    ws["A5"] = 2
    ws["B5"] = 10002
    ws["C5"] = "เด็กหญิงสมหญิง ใจดี"
    wb.save(path)

    res = import_roster(path, conn=migrated_db)
    assert res.ok is True
    assert res.data["students"] >= 2
    assert res.data["enrollments"] >= 2

    cur = migrated_db.cursor()
    cur.execute("SELECT * FROM pragma_foreign_keys")
    fk_on = cur.fetchone()[0]
    assert fk_on == 1, "FK enforcement not ON after successful import!"
