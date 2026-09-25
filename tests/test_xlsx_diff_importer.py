import sqlite3
from unittest.mock import patch
from core.importers.xlsx_diff_importer import apply_xlsx_diff, calculate_xlsx_diff
from data.migrate import apply_migrations


def test_xlsx_diff_categorization_and_atomic_apply():
    conn = sqlite3.connect(":memory:")
    conn.row_factory = sqlite3.Row
    apply_migrations(conn)

    # Seed initial student 101 and existing enrollment in 2566
    conn.execute(
        "INSERT INTO students (student_id, national_id, full_name, first_name, last_name, status) "
        "VALUES (101, '111', 'นายเดิม แท้', 'เดิม', 'แท้', 'active')"
    )
    conn.execute(
        "INSERT INTO enrollments (id, student_id, academic_year, semester, grade_level, room) "
        "VALUES (1, 101, 2566, 2, 'ม.1', 1)"
    )
    conn.commit()

    # Mocked parsed rows from a new yearly xlsx roster
    parsed_rows = [
        # Student 101 promoted from ม.1 room 1 to ม.2 room 1
        {
            "student_id": 101,
            "national_id": "111",
            "prefix": "นาย",
            "first_name": "เดิม",
            "last_name": "แท้",
            "full_name": "นายเดิม แท้",
            "gender": "ชาย",
            "grade_level": "ม.2",
            "room": 1,
            "track": "วิทย์-คณิต",
        },
        # Student 102 newly entering school
        {
            "student_id": 102,
            "national_id": "222",
            "prefix": "เด็กหญิง",
            "first_name": "ใหม่",
            "last_name": "สุข",
            "full_name": "เด็กหญิงใหม่ สุข",
            "gender": "หญิง",
            "grade_level": "ม.1",
            "room": 1,
            "track": None,
        },
    ]

    with patch("core.importers.xlsx_diff_importer.parse_roster_sheets", return_value=parsed_rows):
        res = calculate_xlsx_diff(conn, "dummy.xlsx", academic_year=2567, semester=1)
        assert res.ok
        diff = res.data

        # Verify categorization per design doc §6a
        assert len(diff["new_students"]) == 1
        assert diff["new_students"][0]["student_id"] == 102

        assert len(diff["grade_room_changes"]) == 1
        change = diff["grade_room_changes"][0]
        assert change["student_id"] == 101
        assert change["old_grade"] == "ม.1"
        assert change["new_grade"] == "ม.2"

        # Apply diff atomically
        apply_res = apply_xlsx_diff(conn, diff, academic_year=2567, semester=1, actor="registrar")
        assert apply_res.ok

        # Verify new student inserted
        s102 = conn.execute("SELECT * FROM students WHERE student_id = 102").fetchone()
        assert s102 is not None
        assert s102["full_name"] == "เด็กหญิงใหม่ สุข"

        # Verify student 101 new enrollment period added without clobbering history
        e101 = conn.execute(
            "SELECT * FROM enrollments WHERE student_id = 101 AND academic_year = 2567"
        ).fetchone()
        assert e101 is not None
        assert e101["grade_level"] == "ม.2"

        # Previous enrollment still preserved
        all_e = conn.execute("SELECT * FROM enrollments WHERE student_id = 101").fetchall()
        assert len(all_e) == 2


def test_student_diff_import_and_promotion_real_file(tmp_path):
    """End-to-end real xlsx diff calculation and promotion without mocks."""
    import openpyxl
    conn = sqlite3.connect(":memory:")
    conn.row_factory = sqlite3.Row
    apply_migrations(conn)

    # Insert existing student 1001 in ม.1/1
    conn.execute(
        "INSERT INTO students (student_id, national_id, prefix, first_name, last_name, full_name, status) "
        "VALUES (1001, '1100100200300', 'นาย', 'สมชาย', 'มั่นคง', 'นายสมชาย มั่นคง', 'active')"
    )
    conn.execute(
        "INSERT INTO enrollments (id, student_id, academic_year, semester, grade_level, room) "
        "VALUES (1, 1001, 2567, 1, 'ม.1', 1)"
    )
    conn.commit()

    # Create real multi-sheet roster workbook
    wb = openpyxl.Workbook()
    for g in ["ม.1", "ม.2", "ม.3", "ม.4", "ม.5", "ม.6"]:
        if g not in wb.sheetnames:
            wb.create_sheet(title=g)

    # On sheet ม.2: student 1001 promoted to ม.2/1
    ws_m2 = wb["ม.2"]
    ws_m2.cell(row=1, column=1, value="ชั้นมัธยมศึกษาปีที่ 2/1 ภาคเรียนที่ 1 ปีการศึกษา 2568")
    ws_m2.cell(row=2, column=1, value="ที่")
    ws_m2.cell(row=2, column=2, value="รหัสนักเรียน")
    ws_m2.cell(row=2, column=3, value="ชื่อ - สกุล")
    ws_m2.cell(row=2, column=4, value="เลขประจำตัวประชาชน")
    ws_m2.cell(row=3, column=1, value=1)
    ws_m2.cell(row=3, column=2, value=1001)
    ws_m2.cell(row=3, column=3, value="นายสมชาย มั่นคง")
    ws_m2.cell(row=3, column=4, value="1100100200300")

    # On sheet ม.1: new student 2002 entering ม.1/1
    ws_m1 = wb["ม.1"]
    ws_m1.cell(row=1, column=1, value="ชั้นมัธยมศึกษาปีที่ 1/1 ภาคเรียนที่ 1 ปีการศึกษา 2568")
    ws_m1.cell(row=2, column=1, value="ที่")
    ws_m1.cell(row=2, column=2, value="รหัสนักเรียน")
    ws_m1.cell(row=2, column=3, value="ชื่อ - สกุล")
    ws_m1.cell(row=2, column=4, value="เลขประจำตัวประชาชน")
    ws_m1.cell(row=3, column=1, value=1)
    ws_m1.cell(row=3, column=2, value=2002)
    ws_m1.cell(row=3, column=3, value="เด็กหญิงปราณี สดใส")
    ws_m1.cell(row=3, column=4, value="1100100200999")

    xlsx_path = str(tmp_path / "roster_promotion_2568.xlsx")
    wb.save(xlsx_path)

    # Calculate diff
    diff_res = calculate_xlsx_diff(conn, xlsx_path)
    assert diff_res.ok
    diff = diff_res.data
    assert len(diff["new_students"]) == 1
    assert diff["new_students"][0]["student_id"] == 2002
    assert len(diff["grade_room_changes"]) == 1
    assert diff["grade_room_changes"][0]["student_id"] == 1001
    assert diff["grade_room_changes"][0]["new_grade"] == "ม.2"

    # Apply diff
    apply_res = apply_xlsx_diff(conn, diff, academic_year=2568, semester=1, actor="Admin")
    assert apply_res.ok
    assert apply_res.data["new_added"] == 1
    assert apply_res.data["promoted"] == 1

    # Verify history preserved
    enrs = conn.execute("SELECT * FROM enrollments WHERE student_id = 1001 ORDER BY academic_year ASC").fetchall()
    assert len(enrs) == 2
    assert enrs[0]["grade_level"] == "ม.1"
    assert enrs[0]["academic_year"] == 2567
    assert enrs[1]["grade_level"] == "ม.2"
    assert enrs[1]["academic_year"] == 2568

    # Verify new student active
    s2002 = conn.execute("SELECT * FROM students WHERE student_id = 2002").fetchone()
    assert s2002 is not None
    assert s2002["status"] == "active"
    conn.close()

