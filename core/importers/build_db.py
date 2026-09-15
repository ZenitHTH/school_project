"""
build_db.py — converts raw school-roster xlsx into SQLite tables.
Can be run standalone or imported as a service module.
"""
import sys
import re
import sqlite3
import datetime
from typing import Optional, Dict, Any
from core.shared.result import Result

GRADE_SHEETS = ["ม.1", "ม.2", "ม.3", "ม.4", "ม.5", "ม.6"]
PREFIXES = ["เด็กชาย", "เด็กหญิง", "นางสาว", "นาย", "นาง", "ด.ช.", "ด.ญ."]
MALE_PREFIXES = {"เด็กชาย", "นาย", "ด.ช."}
FEMALE_PREFIXES = {"เด็กหญิง", "นางสาว", "นาง", "ด.ญ."}


def clean(s):
    if s is None:
        return None
    if isinstance(s, str):
        s = s.replace("\xa0", " ").strip()
        return s if s else None
    return s


def split_name(raw_name):
    name = clean(raw_name)
    if not name:
        return None, None, None, None, None
    prefix = None
    rest = name
    for p in sorted(PREFIXES, key=len, reverse=True):
        if name.startswith(p):
            prefix = p
            rest = name[len(p):].strip()
            break
    parts = rest.split()
    first_name = parts[0] if parts else None
    last_name = " ".join(parts[1:]) if len(parts) > 1 else None
    gender = "M" if prefix in MALE_PREFIXES else ("F" if prefix in FEMALE_PREFIXES else None)
    return prefix, first_name, last_name, name, gender


def normalize_date(val):
    if val is None:
        return None
    if isinstance(val, datetime.datetime):
        y = val.year - 543 if val.year > 2400 else val.year
        try:
            return datetime.date(y, val.month, val.day).isoformat()
        except ValueError:
            return f"{y:04d}-{val.month:02d}-{val.day:02d}"
    s = clean(val)
    if not s:
        return None
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})$", s)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y > 2400:
            y -= 543
        try:
            return datetime.date(y, mo, d).isoformat()
        except ValueError:
            return s
    return s


def to_int(val):
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return int(val)
    s = clean(val)
    if s and s.isdigit():
        return int(s)
    return None


def upsert_student(cur, student_id, national_id, raw_name):
    if student_id is None or not raw_name:
        return
    prefix, first, last, full, gender = split_name(raw_name)
    cur.execute(
        """
        INSERT INTO students (student_id, national_id, prefix, first_name, last_name, full_name, gender)
        VALUES (?,?,?,?,?,?,?)
        ON CONFLICT(student_id) DO UPDATE SET
            national_id=COALESCE(excluded.national_id, students.national_id),
            prefix=excluded.prefix, first_name=excluded.first_name,
            last_name=excluded.last_name, full_name=excluded.full_name,
            gender=excluded.gender
        """,
        (student_id, clean(str(national_id)) if national_id is not None else None,
         prefix, first, last, full, gender),
    )


TITLE_RE = re.compile(r"ชั้นมัธยมศึกษาปีที่\s*(\d+)\s*/\s*(\d+)")
YEAR_RE = re.compile(r"ปีการศึกษา\s*(\d{4})")
SEM_RE = re.compile(r"ภาคเรียนที่\s*(\d+)")


def parse_grade_sheet(ws, cur):
    grade, room, year, semester, track = None, None, None, None, None
    in_block = False
    for r in range(1, ws.max_row + 1):
        a = clean(ws.cell(row=r, column=1).value)
        b = ws.cell(row=r, column=2).value
        c = clean(ws.cell(row=r, column=3).value)
        d = ws.cell(row=r, column=4).value

        title_match = TITLE_RE.search(str(a)) if a else None
        if title_match:
            grade, room = int(title_match.group(1)), int(title_match.group(2))
            ym = YEAR_RE.search(a)
            sm = SEM_RE.search(a)
            year = int(ym.group(1)) - 543 if ym else None
            semester = int(sm.group(1)) if sm else None
            track = TITLE_RE.sub("", a)
            track = YEAR_RE.sub("", track)
            track = SEM_RE.sub("", track)
            track = clean(track.replace("ภาคเรียนที่", "")) or None
            in_block = False
            continue
        if a == "ที่":
            in_block = True
            continue
        if in_block and isinstance(b, (int, float)) and c:
            seq = to_int(a)
            student_id = to_int(b)
            upsert_student(cur, student_id, d, c)
            if student_id and grade and room:
                cur.execute(
                    """INSERT OR IGNORE INTO enrollments
                       (student_id, academic_year, semester, grade_level, room, track, seq_no)
                       VALUES (?,?,?,?,?,?,?)""",
                    (student_id, year, semester, f"ม.{grade}", room, track, seq),
                )
        elif in_block and a is None and b is None:
            in_block = False


def parse_transfers_out(ws, cur, source_sheet, header_rows):
    for r in range(header_rows + 1, ws.max_row + 1):
        student_id = to_int(ws.cell(row=r, column=2).value)
        name = ws.cell(row=r, column=3).value
        room = clean(ws.cell(row=r, column=5).value)
        dest = clean(ws.cell(row=r, column=6).value)
        date = normalize_date(ws.cell(row=r, column=7).value)
        note = clean(ws.cell(row=r, column=8).value)
        if not student_id or not name:
            continue
        upsert_student(cur, student_id, ws.cell(row=r, column=4).value, name)
        cur.execute(
            """INSERT INTO transfers_out (student_id, room_at_transfer, destination_school, transfer_date, note, source_sheet)
               VALUES (?,?,?,?,?,?)""",
            (student_id, room, dest, date, note, source_sheet),
        )
        cur.execute("UPDATE students SET status='transferred_out' WHERE student_id=?", (student_id,))


def parse_transfers_in(ws, cur):
    for r in range(4, ws.max_row + 1):
        old_id = to_int(ws.cell(row=r, column=2).value)
        new_id = to_int(ws.cell(row=r, column=3).value)
        name = ws.cell(row=r, column=4).value
        prev_school = clean(ws.cell(row=r, column=5).value)
        date = normalize_date(ws.cell(row=r, column=6).value)
        room = clean(ws.cell(row=r, column=7).value)
        note = clean(ws.cell(row=r, column=8).value)
        student_id = new_id or old_id
        if not student_id or not name:
            continue
        upsert_student(cur, student_id, None, name)
        cur.execute(
            """INSERT INTO transfers_in (student_id, old_student_id, previous_school, transfer_date, room, note)
               VALUES (?,?,?,?,?,?)""",
            (student_id, old_id, prev_school, date, room, note),
        )


def parse_room_changes(ws, cur):
    for r in range(5, ws.max_row + 1):
        student_id = to_int(ws.cell(row=r, column=2).value)
        name = ws.cell(row=r, column=3).value
        old_room = clean(ws.cell(row=r, column=4).value)
        new_room = clean(ws.cell(row=r, column=5).value)
        date = normalize_date(ws.cell(row=r, column=6).value)
        note = clean(ws.cell(row=r, column=7).value)
        if not student_id or not name:
            continue
        upsert_student(cur, student_id, None, name)
        cur.execute(
            """INSERT INTO room_changes (student_id, old_room, new_room, change_date, note)
               VALUES (?,?,?,?,?)""",
            (student_id, old_room, new_room, date, note),
        )


def parse_leave(ws, cur):
    for r in range(4, ws.max_row + 1):
        student_id = to_int(ws.cell(row=r, column=2).value)
        name = ws.cell(row=r, column=3).value
        national_id = ws.cell(row=r, column=4).value
        grade_room = clean(ws.cell(row=r, column=5).value)
        leave_date = normalize_date(ws.cell(row=r, column=6).value)
        note = clean(ws.cell(row=r, column=7).value)
        return_date = normalize_date(ws.cell(row=r, column=8).value)
        if not student_id or not name:
            continue
        upsert_student(cur, student_id, national_id, name)
        cur.execute(
            """INSERT INTO leave_of_absence (student_id, grade_room_at_leave, leave_date, expected_return_date, note)
               VALUES (?,?,?,?,?)""",
            (student_id, grade_room, leave_date, return_date, note),
        )
        cur.execute("UPDATE students SET status='on_leave' WHERE student_id=?", (student_id,))


def parse_retentions_69(ws, cur, year):
    for r in range(2, ws.max_row + 1):
        student_id = to_int(ws.cell(row=r, column=2).value)
        name = ws.cell(row=r, column=3).value
        grade = clean(ws.cell(row=r, column=4).value)
        entry_date = normalize_date(ws.cell(row=r, column=5).value)
        note = clean(ws.cell(row=r, column=6).value)
        if not student_id or not name:
            continue
        upsert_student(cur, student_id, None, name)
        cur.execute(
            """INSERT INTO retentions (student_id, academic_year, new_class, entry_date, note)
               VALUES (?,?,?,?,?)""",
            (student_id, year, grade, entry_date, note),
        )


def parse_retentions_68(ws, cur, year):
    for r in range(3, ws.max_row + 1):
        student_id = to_int(ws.cell(row=r, column=2).value)
        name = ws.cell(row=r, column=3).value
        old_class = clean(ws.cell(row=r, column=5).value)
        new_class = clean(ws.cell(row=r, column=6).value)
        note = clean(ws.cell(row=r, column=7).value)
        if not student_id or not name:
            continue
        upsert_student(cur, student_id, None, name)
        cur.execute(
            """INSERT INTO retentions (student_id, academic_year, old_class, new_class, note)
               VALUES (?,?,?,?,?)""",
            (student_id, year, old_class, new_class, note),
        )


def parse_repeat_room(ws, cur):
    for r in range(2, ws.max_row + 1):
        student_id = to_int(ws.cell(row=r, column=2).value)
        name = ws.cell(row=r, column=3).value
        entry_date = normalize_date(ws.cell(row=r, column=4).value)
        room = clean(ws.cell(row=r, column=5).value)
        note = clean(ws.cell(row=r, column=6).value)
        if not student_id or not name:
            continue
        upsert_student(cur, student_id, None, name)
        cur.execute(
            """INSERT INTO retentions (student_id, new_class, entry_date, note)
               VALUES (?,?,?,?)""",
            (student_id, room, entry_date, note),
        )


def import_roster(xlsx_path: str, conn: Optional[sqlite3.Connection] = None, db_path: Optional[str] = None) -> Result:
    """Import roster data from xlsx into open connection or db_path."""
    try:
        import openpyxl
    except ImportError:
        return Result.fail("openpyxl library is required for importing xlsx rosters")

    close_after = False
    if conn is None:
        if not db_path:
            return Result.fail("Must provide either an active database connection or db_path")
        conn = sqlite3.connect(db_path)
        close_after = True

    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        cur = conn.cursor()

        # Turn foreign keys off during bulk roster load
        conn.execute("PRAGMA foreign_keys = OFF;")
        try:
            with conn:
                # Clear student-domain tables only (never drop library tables or activity_log!)
                for t in ["retentions", "leave_of_absence", "transfers_in", "transfers_out", "room_changes", "enrollments", "students"]:
                    cur.execute(f"DELETE FROM {t};")

                for name in GRADE_SHEETS:
                    if name in wb.sheetnames:
                        parse_grade_sheet(wb[name], cur)

                if "ย้ายออก" in wb.sheetnames:
                    parse_transfers_out(wb["ย้ายออก"], cur, "ย้ายออก", header_rows=1)
                if "ย้ายออกส่งเขต1" in wb.sheetnames:
                    parse_transfers_out(wb["ย้ายออกส่งเขต1"], cur, "ย้ายออกส่งเขต1", header_rows=3)
                if "ย้ายออกส่งเขต (2)" in wb.sheetnames:
                    parse_transfers_out(wb["ย้ายออกส่งเขต (2)"], cur, "ย้ายออกส่งเขต (2)", header_rows=3)
                if "ย้ายเข้า" in wb.sheetnames:
                    parse_transfers_in(wb["ย้ายเข้า"], cur)
                if "ย้ายห้อง" in wb.sheetnames:
                    parse_room_changes(wb["ย้ายห้อง"], cur)
                if "รายชื่อพักการเรียน" in wb.sheetnames:
                    parse_leave(wb["รายชื่อพักการเรียน"], cur)
                if "รายชื่อไม่ได้เลื่อนชั้นปี69" in wb.sheetnames:
                    parse_retentions_69(wb["รายชื่อไม่ได้เลื่อนชั้นปี69"], cur, 2026)
                if "รายชื่อไม่ได้เลื่อนชั้นปี68" in wb.sheetnames:
                    parse_retentions_68(wb["รายชื่อไม่ได้เลื่อนชั้นปี68"], cur, 2025)
                if "นักเรียนม.3ชั้น" in wb.sheetnames:
                    parse_repeat_room(wb["นักเรียนม.3ชั้น"], cur)

            counts = {}
            for t in ["students", "enrollments", "room_changes", "transfers_out", "transfers_in", "leave_of_absence", "retentions"]:
                n = cur.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
                counts[t] = n

            return Result.success(counts)
        except Exception as e:
            return Result.fail(f"Failed to import roster: {e}")
        finally:
            # Always restore FK enforcement regardless of success or failure
            conn.execute("PRAGMA foreign_keys = ON;")
    finally:
        if close_after:
            conn.close()


if __name__ == "__main__":
    import sys
    x_path = sys.argv[1] if len(sys.argv) > 1 else "default.profraw.xlsx"
    d_path = sys.argv[2] if len(sys.argv) > 2 else "school.sqlite"
    res = import_roster(x_path, db_path=d_path)
    if res.ok:
        print("Import successful:")
        for tbl, count in res.data.items():
            print(f"  {tbl}: {count} rows")
    else:
        print(f"Import failed: {res.error}")
